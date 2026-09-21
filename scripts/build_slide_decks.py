"""Generate PowerPoint decks from the five slide plans in presentations/.

Each plan already specifies, per slide, the visible copy, the intended visual and
the speaker notes. This turns those specifications into editable 16:9 decks that
follow the shared presentation brief: white background, dark navy text, teal
accent, large type, one idea per slide.

    python -m pip install python-pptx
    python scripts/build_slide_decks.py

The visual description is carried onto the slide as a placeholder note so that
whoever builds the chart or diagram can see what belongs there, and it is also
repeated at the top of the speaker notes. Replace each placeholder with the real
chart or diagram before delivery, and check every figure against the plan.
"""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
PLANS = sorted((ROOT / "presentations").glob("day-*-slide-plan.md"))
OUTPUT = ROOT / "dist" / "slides"

NAVY = RGBColor(0x14, 0x21, 0x3D)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
GREY = RGBColor(0x5A, 0x64, 0x78)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.9)

SLIDE_HEADING = re.compile(r"^##\s+Slide\s+(\d+)\s*:\s*(.*)$")
FIELD = re.compile(r"^(Visible copy|Visual|Speaker notes)\s*:\s*(.*)$")


class Slide:
    def __init__(self, number: int, title: str) -> None:
        self.number = number
        self.title = title
        self.copy: list[str] = []
        self.visual = ""
        self.notes = ""

    @property
    def is_cover(self) -> bool:
        return self.number == 1

    @property
    def is_break(self) -> bool:
        return self.title.strip().lower() == "break"


def parse(path: Path) -> tuple[str, list[Slide]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    deck_title = next((l[2:].strip() for l in lines if l.startswith("# ")), path.stem)

    slides: list[Slide] = []
    current: Slide | None = None
    field: str | None = None

    for line in lines:
        heading = SLIDE_HEADING.match(line.strip())
        if heading:
            current = Slide(int(heading.group(1)), heading.group(2).strip())
            slides.append(current)
            field = None
            continue
        if current is None:
            continue

        match = FIELD.match(line.strip())
        if match:
            field, rest = match.group(1), match.group(2).strip()
            if field == "Visual":
                current.visual = rest
            elif field == "Speaker notes":
                current.notes = rest
            elif rest:
                current.copy.append(rest)
            continue

        text = line.strip()
        if not text:
            continue
        if field == "Visible copy":
            current.copy.append(text)
        elif field == "Visual":
            current.visual = f"{current.visual} {text}".strip()
        elif field == "Speaker notes":
            current.notes = f"{current.notes} {text}".strip()

    return deck_title, slides


def textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    return frame


def write(frame, text, *, size, colour, bold=False, space_after=10, align=PP_ALIGN.LEFT, first=False):
    paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
    paragraph.alignment = align
    paragraph.space_after = Pt(space_after)
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = colour
    run.font.name = "Arial"
    return paragraph


def accent_rule(slide, left, top, width):
    from pptx.enum.shapes import MSO_SHAPE

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Emu(38100))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def build(deck_title: str, slides: list[Slide], destination: Path) -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    for spec in slides:
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PAPER

        if spec.is_cover:
            frame = textbox(slide, MARGIN, Inches(2.4), SLIDE_W - 2 * MARGIN, Inches(3.0))
            write(frame, spec.title, size=46, colour=NAVY, bold=True, space_after=18, first=True)
            for entry in spec.copy:
                write(frame, entry, size=26, colour=GREY, space_after=6)
            accent_rule(slide, MARGIN, Inches(2.15), Inches(3.2))

        elif spec.is_break:
            frame = textbox(slide, MARGIN, Inches(3.0), SLIDE_W - 2 * MARGIN, Inches(1.6))
            write(frame, spec.title, size=40, colour=GREY, bold=True,
                  align=PP_ALIGN.CENTER, space_after=14, first=True)
            for entry in spec.copy:
                write(frame, entry, size=30, colour=NAVY, align=PP_ALIGN.CENTER, space_after=6)

        else:
            title_frame = textbox(slide, MARGIN, Inches(0.55), SLIDE_W - 2 * MARGIN, Inches(1.0))
            write(title_frame, spec.title, size=34, colour=NAVY, bold=True, space_after=0, first=True)
            accent_rule(slide, MARGIN, Inches(1.45), Inches(1.8))

            body = textbox(slide, MARGIN, Inches(1.85), SLIDE_W - 2 * MARGIN, Inches(3.4))
            for index, entry in enumerate(spec.copy):
                write(body, entry, size=26, colour=NAVY, space_after=14, first=(index == 0))

            if spec.visual:
                hint = textbox(slide, MARGIN, Inches(5.65), SLIDE_W - 2 * MARGIN, Inches(1.2))
                write(hint, f"[ VISUAL TO ADD ]  {spec.visual}", size=13, colour=TEAL,
                      space_after=0, first=True)

        parts = [f"Slide {spec.number}"]
        if spec.visual:
            parts.append(f"Visual: {spec.visual}")
        if spec.notes:
            parts.append(spec.notes)
        parts.append("Pause for Arabic interpretation before moving on.")
        slide.notes_slide.notes_text_frame.text = "\n\n".join(parts)

    destination.parent.mkdir(parents=True, exist_ok=True)
    prs.save(destination)


def main() -> None:
    if not PLANS:
        raise SystemExit("no slide plans found in presentations/")
    for plan in PLANS:
        deck_title, slides = parse(plan)
        if not slides:
            print(f"skipped {plan.name}: no slides parsed")
            continue
        destination = OUTPUT / f"{plan.stem.replace('-slide-plan', '')}.pptx"
        build(deck_title, slides, destination)
        placeholders = sum(1 for s in slides if s.visual and not s.is_cover and not s.is_break)
        print(f"{destination.relative_to(ROOT)}: {len(slides)} slides, "
              f"{placeholders} visual placeholders  ({deck_title})")


if __name__ == "__main__":
    main()
